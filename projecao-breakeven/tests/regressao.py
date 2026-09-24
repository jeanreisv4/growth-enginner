# -*- coding: utf-8 -*-
"""Regressão da skill: roda piloto + gerador nos dados sintéticos e confere as regras que o usuário exigiu.

  python3 tests/regressao.py

Casos: inside sales, e-commerce sem GA4 e e-commerce com GA4 (Google e Meta separados), mais a planilha demo.
Confere em cada planilha, recalculando as fórmulas com pycel:
  1. nenhuma fórmula com erro;
  2. nenhuma célula azul;
  3. a receita projetada no template é a mesma do piloto, mês a mês (o template não pode inflar nem perder receita);
  4. nenhuma taxa de etapa passa de 100% (connect, CTR, cliques → sessões, sessão → view item...);
  5. o número de meses é o horizonte pedido.
Precisa de pandas, openpyxl e pycel (python3 -m pip install --user pandas openpyxl pycel).
"""
import json, os, subprocess, sys, tempfile, warnings
warnings.filterwarnings("ignore")
import openpyxl
from pycel import ExcelCompiler

SKILL = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIX = os.path.join(SKILL, "tests", "fixtures")
PILOTO = os.path.join(SKILL, "scripts", "breakeven_pilot.py")
GERADOR = os.path.join(SKILL, "scripts", "gerar_template.py")
AZUL = {"FF0000FF", "000000FF", "FF4472C4", "4472C4", "FF1F4E79", "FF2E75B6", "FF5B9BD5", "FF0070C0"}
HORIZONTE = 4
CASOS = [
    dict(nome="inside_sales", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=[], extra_gerador=[], exige=["[%] CONEXÃO SOBRE OS LEADS (MQL OU NÃO)", "[QNTD] CONEXÕES", "[%] CONEXÃO → SQL"],
         sem=["[QNTD] MQLS CONECTADOS", "[%] CONEXÃO — MQL ATENDIDO"]),
    dict(nome="inside_sales_irrealista", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", fee="40000", extra_piloto=["--fixar", "clique_lead=0.05"], extra_gerador=[],
         metodologia=["CAMINHO PARA O BREAKEVEN"], exige=["[QNTD] CONEXÕES", "[%] CLIQUE → LEAD"],
         sem=["[QNTD] VISITAS TOTAIS NA PÁGINA", "[%] CONNECT RATE (CLIQUE → VISITA NA PÁGINA)", "[QNTD] MQLS CONECTADOS"]),
    dict(nome="inside_sales_sem_conexao", modelo="inside_sales", csv="indicadores_inside_sales_sem_conexao.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=[], extra_gerador=[],
         exige=["[%] LEAD → MQL", "[QNTD] MQLS", "[%] MQL → SQL"],
         sem=["[%] CONEXÃO SOBRE OS LEADS (MQL OU NÃO)", "[QNTD] CONEXÕES", "[%] CONEXÃO → SQL"]),
    # cliente em dólar: --moeda tem de trocar formato numérico E texto do rótulo, sem encostar em fórmula
    # ("R$" também é referência absoluta válida: coluna R, linha travada).
    dict(nome="inside_sales_moeda_usd", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[US$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=[], extra_gerador=["--moeda", "US$"],
         exige=["[US$] FEE V4", "[US$] CUSTO TOTAL (FEE V4 + MÍDIA)", "[US$] TICKET MÉDIO"],
         sem=["[R$] FEE V4", "[R$] FATURAMENTO (VENDAS × TICKET)", "[R$] TICKET MÉDIO"]),
    dict(nome="inside_sales_lp", modelo="inside_sales", csv="indicadores_inside_sales_lp.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=[], extra_gerador=[],
         exige=["[%] CONNECT RATE (CLIQUE → VISITA NA PÁGINA)", "[QNTD] VISITAS PAGAS NA PÁGINA", "[R$] CUSTO POR VISITA", "[%] VISITA → LEAD"],
         sem=["[%] CLIQUE → LEAD", "[QNTD] VISITAS TOTAIS NA PÁGINA"],
         valores={"[%] CONNECT RATE (CLIQUE → VISITA NA PÁGINA)": 0.80, "[%] VISITA → LEAD": 0.10}),
    dict(nome="inside_sales_connect_informado", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=["--connect-rate", "0.8"], extra_gerador=[],
         exige=["[%] CONNECT RATE (CLIQUE → VISITA NA PÁGINA)", "[QNTD] VISITAS PAGAS NA PÁGINA", "[%] VISITA → LEAD"],
         sem=["[%] CLIQUE → LEAD"],
         valores={"[%] CONNECT RATE (CLIQUE → VISITA NA PÁGINA)": 0.80, "[%] VISITA → LEAD": 0.10}),
    dict(nome="inside_sales_conexao_so_lead", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=[], extra_gerador=["--conexao-so-lead"],
         exige=["[QNTD] CONEXÕES", "[%] CONEXÃO → SQL"],
         sem=["[%] COMISSÃO SOBRE O GMV", "[R$] RECEITA DA AGÊNCIA (GMV × COMISSÃO)", "[QNTD] MQLS CONECTADOS"]),
    dict(nome="inside_sales_mercado", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=["--alvo", "conexao=0.9", "--alvo", "sql_venda=0.5", "--verba-plano", "5000,6000,7000,8000"],
         extra_gerador=["--conexao-so-lead"], exige=["[%] CONEXÃO SOBRE OS LEADS (MQL OU NÃO)", "[QNTD] CONEXÕES"]),
    dict(nome="inside_sales_sazonalidade", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=["--sazonalidade-demanda", "1,1,1.3", "--sazonalidade-cpm", "1,1,1.2", "--cpm-crescimento", "0.04", "--cpm-crescimento-ate", "3"], extra_gerador=[],
         metodologia=["CPM MÊS A MÊS", "SAZONALIDADE DA DEMANDA MÊS A MÊS"], exige=["[R$] CPM (CUSTO POR MIL IMPRESSÕES)", "[%] SQL → VENDA", "[X] ROAS (FATURAMENTO ÷ MÍDIA)",
                "[R$] RETORNO DOS ANÚNCIOS APÓS A MARGEM (MC − MÍDIA)", "[R$] FATURAMENTO NECESSÁRIO PARA ZERAR O MÊS",
                "[QNTD] VENDAS NECESSÁRIAS PARA ZERAR O MÊS", "[X] RETORNO POR R$ 1 INVESTIDO NO MÊS (MC ÷ CUSTO)"],
         sem=["MULTIPLICADOR DO CPM (ELEIÇÃO, DATAS, REMARKETING)", "CPM EFETIVO (BASE × MULTIPLICADOR)", "SAZONALIDADE DA DEMANDA (× VENDAS)"]),
    dict(nome="inside_sales_extra", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=[], extra_gerador=["--extra", "premissas_inside_sales_extra.json|Cenário B|Cenário alternativo"],
         abas=["Cenário B", "Premissas · Cenário B"]),
    dict(nome="inside_sales_seo", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=["--organico-visitas", "0,500,1000,2000", "--organico-conversao", "0.05"], extra_gerador=[],
         exige=["[QNTD] VISITAS ORGÂNICAS (SEO)", "[%] VISITAS ORGÂNICAS → LEAD", "[QNTD] LEADS ORGÂNICOS (SEO)",
                "[QNTD] LEADS TOTAIS (PAGOS + ORGÂNICOS)", "[QNTD] CONEXÕES"]),
    dict(nome="inside_sales_social", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_gerador=[],
         extra_piloto=["--organico-visitas", "0,300,600,900", "--organico-conversao", "0.04",
                       "--organico-origem", "social orgânico", "--organico-metrica", "cliques no link da bio"],
         exige=["[QNTD] CLIQUES NO LINK DA BIO (SOCIAL ORGÂNICO)", "[%] CLIQUES NO LINK DA BIO → LEAD",
                "[QNTD] LEADS ORGÂNICOS (SOCIAL ORGÂNICO)", "[QNTD] LEADS TOTAIS (PAGOS + ORGÂNICOS)"],
         sem=["[QNTD] VISITAS ORGÂNICAS (SEO)"]),
    dict(nome="inside_sales_crm", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[1] GMV · NOVOS LEADS (VENDAS × TICKET)",
         extra_piloto=["--crm", os.path.join(FIX, "crm_inside_sales.json"), "--fee-plano", "1200,4200", "--organico-visitas", "0,500,1000,2000", "--organico-conversao", "0.05"],
         extra_gerador=[], series={"[2] GMV · RECOMPRA DA BASE (WHATSAPP + CROSS-SELL)": "receita_recompra",
                                   "[3] GMV · REATIVAÇÃO (E-MAIL PARA LEADS + INATIVOS)": "receita_reativacao"},
         exige=["[QNTD] LEADS SEM COMPRA NA BASE (INÍCIO DO MÊS)", "[R$] GMV TOTAL (NOVOS + RECOMPRA + REATIVAÇÃO)"]),
    dict(nome="inside_sales_legado", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=[],
         extra_gerador=["--legado", os.path.join(FIX, "legado_inside_sales.json")],
         metodologia=["PROJETADO × REALIZADO"]),
    dict(nome="inside_sales_ciclo", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         receita="[R$] FATURAMENTO (VENDAS × TICKET)", extra_piloto=["--lag", "0.5"], extra_gerador=[],
         exige=["[QNTD] VENDAS ORIGINADAS PELOS SQLS DO MÊS", "[QNTD] VENDAS FECHADAS NO MÊS"]),
    dict(nome="inside_sales_comissao", modelo="inside_sales", csv="indicadores_inside_sales.csv", aba="Inside Sales",
         # a linha "receita" do piloto é o topo de funil (vendas × ticket); a comissão só entra depois, na receita da agência
         margem="1", comissao="0.15", receita="[R$] GMV (VENDAS × TICKET)", extra_piloto=[], extra_gerador=[],
         exige=["[%] COMISSÃO SOBRE O GMV (MARGEM DA AGÊNCIA)", "[R$] RESULTADO MC (= RECEITA DA AGÊNCIA)",
                "[X] ROAS DE BREAKEVEN (RECEITA NECESSÁRIA ÷ MÍDIA)", "[R$] RECEITA NECESSÁRIA PARA ZERAR O MÊS"],
         sem=["[%] MARGEM DE CONTRIBUIÇÃO"]),
    dict(nome="ecommerce", modelo="ecommerce", csv="indicadores_ecommerce.csv", aba="E-commerce",
         receita="[R$] RECEITA ATRIBUÍDA À MÍDIA V4", extra_piloto=[], extra_gerador=["--sem-etapa-venda"]),
    dict(nome="ecommerce_ga4", modelo="ecommerce", csv="indicadores_ecommerce.csv", aba="E-commerce",
         receita="[R$] RECEITA ATRIBUÍDA À MÍDIA V4", extra_piloto=["--ga4", os.path.join(FIX, "ga4_ecommerce.json")],
         extra_gerador=["--sem-etapa-venda"], exige=["[R$] VERBA GOOGLE", "[R$] VERBA META", "[QNTD] SESSÕES - NÃO PAGAS"]),
]


def rodar(cmd, cwd):
    r = subprocess.run([sys.executable] + cmd, cwd=cwd, capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(f"falhou: {' '.join(cmd[:2])}\n{r.stderr[-1500:]}")
    return r.stdout


def conferir_planilha(path, aba, rot_receita=None, projecao=None, exige=(), metodologia=(), sem=(), series=None, valores=None):
    wb = openpyxl.load_workbook(path); exc = ExcelCompiler(filename=path); falhas = []
    for nome in wb.sheetnames:
        ws = wb[nome]
        for row in ws.iter_rows():
            for c in row:
                if c.fill is not None and c.fill.fgColor is not None and str(c.fill.fgColor.rgb) in AZUL:
                    falhas.append(f"{nome}!{c.coordinate}: célula azul")
                if isinstance(c.value, str) and c.value.startswith("="):
                    try:
                        v = exc.evaluate(f"'{nome}'!{c.coordinate}")
                    except Exception as e:
                        falhas.append(f"{nome}!{c.coordinate}: {type(e).__name__}"); continue
                    if isinstance(v, str) and v.startswith("#"):
                        falhas.append(f"{nome}!{c.coordinate}: {v}")
    ws = wb[aba]
    linhas = {ws.cell(r, 2).value.strip(): r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 2).value, str)}
    cabec = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 3).value == "Projetado" or ws.cell(r, 4).value == "Realizado")
    def e_total(c):  # o par de totais também tem cabeçalho Projetado | Realizado
        acima = ws.cell(cabec - 1, c).value
        if acima is None:  # cabeçalho mesclado: o texto fica na célula à esquerda
            acima = ws.cell(cabec - 1, c - 1).value
        return isinstance(acima, str) and "TOTAL" in acima.upper()
    cols = [ws.cell(cabec, c).column_letter for c in range(3, ws.max_column + 1)
            if ws.cell(cabec, c).value == "Projetado" and not e_total(c)]
    for rot in exige:
        if rot not in linhas: falhas.append(f"linha ausente: {rot}")
    for rot in sem:
        if rot in linhas: falhas.append(f"linha que não devia aparecer: {rot}")
    if metodologia:
        met = next((wb[n] for n in wb.sheetnames if n.startswith("Premissas")), None)
        textos = {c.value for row in met.iter_rows() for c in row if isinstance(c.value, str)} if met else set()
        for t in metodologia:
            if not any(x.startswith(t) for x in textos): falhas.append(f"Premissas sem a seção {t}")
    for rot, r in linhas.items():
        if r <= cabec: continue  # só a tabela de projeção; o topo tem blocos de premissas lado a lado
        up = rot.upper()
        # taxa de etapa: percentual que, por definição, não passa de 100%. Ficam de fora o crescimento mês a mês (pode passar)
        # e a conexão, cuja coluna Realizado copia a fonte, onde "Conexões" às vezes conta tentativas e supera os leads
        # (o piloto trava a taxa projetada em 100% e emite alerta; falsear o realizado seria pior).
        if (rot.startswith("[%]") and not any(x in up for x in ("CRESCIMENTO", "CONEXÃO", "PARTICIPAÇÃO"))) or "→" in rot or "CONNECT" in up or "CTR" in up:
            for c in range(3, 3 + 2 * len(cols)):
                v = exc.evaluate(f"'{aba}'!{ws.cell(r, c).coordinate}")
                if isinstance(v, (int, float)) and v > 1.0001:
                    falhas.append(f"{rot} = {v:.1%} em {ws.cell(r, c).coordinate}: taxa acima de 100%")
    for rot, chave in (series or {}).items():   # outras linhas que precisam bater com o piloto (ex.: GMV do CRM)
        r = linhas.get(rot)
        if r is None:
            falhas.append(f"linha ausente: {rot}"); continue
        for k, (col, lin) in enumerate(zip(cols, projecao or [])):
            v = exc.evaluate(f"'{aba}'!{col}{r}") or 0
            if abs(v - lin.get(chave, 0)) > max(1.0, 0.005 * abs(lin.get(chave, 0))):
                falhas.append(f"M{k + 1}: {rot[:30]} do template {v:,.2f} ≠ piloto {lin.get(chave, 0):,.2f}")
    for rot, esperado in (valores or {}).items():   # prende o número da alavanca, não só a existência da linha
        r = linhas.get(rot)
        if r is None:
            falhas.append(f"linha ausente: {rot}"); continue
        v = exc.evaluate(f"'{aba}'!{cols[0]}{r}")
        if not isinstance(v, (int, float)) or abs(v - esperado) > max(1e-4, 0.01 * abs(esperado)):
            falhas.append(f"{rot}: M1 = {v}, esperado {esperado}")
    # identidade do breakeven do mês: sobra/falta × margem implícita = resultado líquido do mês.
    # É o que garante que a linha "necessária para zerar" zera mesmo o mês, em qualquer variante de receita.
    be = next((r for k, r in linhas.items() if "PARA ZERAR O MÊS" in k and k.startswith("[R$]")), None)
    if be:   # mês vivido sem receita não pode zerar a linha: a margem implícita é 0/0 e tem de cair para a do plano
        for col in cols:
            v = exc.evaluate(f"'{aba}'!{col}{be}")
            if not isinstance(v, (int, float)) or v <= 0:
                falhas.append(f"{col}: receita necessária para zerar o mês = {v} (deveria ser sempre positiva)")
    gap = next((r for k, r in linhas.items() if "SOBRA / FALTA" in k), None)
    res = next((r for k, r in linhas.items() if k.startswith("[R$] RESULTADO LÍQUIDO DO MÊS")), None)
    rec = next((r for k, r in linhas.items() if k in ("[R$] RECEITA CONSOLIDADA", "[R$] FATURAMENTO CONSOLIDADO")), None)
    mc = next((r for k, r in linhas.items() if "RESULTADO MC CONSOLIDADO" in k), None)
    if be and gap and res and rec:
        for col in cols:
            ev = lambda rr: exc.evaluate(f"'{aba}'!{col}{rr}") or 0
            g, r_, f_ = ev(gap), ev(rec), ev(be)
            if abs((r_ - f_) - g) > max(1.0, 0.005 * abs(g)):
                falhas.append(f"{col}: sobra/falta {g:,.2f} ≠ receita − necessário {(r_ - f_):,.2f}")
            m = (ev(mc) / r_) if (mc and r_) else 1.0
            if abs(g * m - ev(res)) > max(1.0, 0.005 * abs(ev(res))):
                falhas.append(f"{col}: sobra/falta × margem {g * m:,.2f} ≠ resultado do mês {ev(res):,.2f}")
    if projecao is not None:
        if len(cols) != len(projecao):
            falhas.append(f"template tem {len(cols)} meses, piloto tem {len(projecao)}")
        r = linhas.get(rot_receita)
        if r is None:
            falhas.append(f"linha de receita não encontrada: {rot_receita}")
        else:
            for k, (col, lin) in enumerate(zip(cols, projecao)):
                v = exc.evaluate(f"'{aba}'!{col}{r}") or 0
                if abs(v - lin["receita"]) > max(1.0, 0.005 * abs(lin["receita"])):
                    falhas.append(f"M{k + 1}: receita do template R$ {v:,.2f} ≠ piloto R$ {lin['receita']:,.2f}")
    return falhas


def main():
    tmp = tempfile.mkdtemp(prefix="regressao_breakeven_")
    total = 0
    for caso in CASOS:
        prem = os.path.join(tmp, f"premissas_{caso['nome']}.json"); xlsx = os.path.join(tmp, f"{caso['nome']}.xlsx")
        try:
            rodar([PILOTO, "projetar", "--fonte", os.path.join(FIX, caso["csv"]), "--aba", "Indicadores", "--modelo", caso["modelo"],
                   "--fee", caso.get("fee", "1200"), "--midia", "5000", "--margem", caso.get("margem", "0.3"),
                   "--comissao", caso.get("comissao", "1"), "--mes-alvo", "3",
                   "--horizonte", str(HORIZONTE), "--crescimento-midia", "0.05", "--out", prem] + caso["extra_piloto"], tmp)
            rodar([GERADOR, "--premissas", prem, "--modelo", caso["modelo"], "--cliente", "Teste", "--out", xlsx] + caso["extra_gerador"], tmp)
            d = json.load(open(prem, encoding="utf-8"))
            if d.get("teto_receita") is None:
                raise RuntimeError("premissas.json sem teto_receita: a trava de sanidade não chegou ao gerador")
            falhas = conferir_planilha(xlsx, caso["aba"], caso["receita"], d["projecao"], caso.get("exige", ()), caso.get("metodologia", ()),
                                       caso.get("sem", ()), caso.get("series"), caso.get("valores"))
            if caso.get("abas"):   # cenário extra em outra aba do mesmo arquivo
                nomes = openpyxl.load_workbook(xlsx).sheetnames
                falhas += [f"aba ausente: {a}" for a in caso["abas"] if a not in nomes]
                if "Cenário B" in nomes: falhas += conferir_planilha(xlsx, "Cenário B", caso["receita"], d["projecao"])
        except RuntimeError as e:
            falhas = [str(e)]
        total += len(falhas)
        print(f"{'OK   ' if not falhas else 'FALHA'} {caso['nome']}" + "".join(f"\n      - {f}" for f in falhas[:12]))
    demo = os.path.join(tmp, "demo.xlsx")
    try:
        rodar([os.path.join(SKILL, "gerador", "build_workbook.py"), demo], tmp)
        falhas = conferir_planilha(demo, "E-commerce") + conferir_planilha(demo, "Inside Sales")
    except RuntimeError as e:
        falhas = [str(e)]
    total += len(falhas)
    print(f"{'OK   ' if not falhas else 'FALHA'} demo" + "".join(f"\n      - {f}" for f in falhas[:12]))
    print(f"\n{'Tudo certo.' if not total else f'{total} falha(s).'}  Arquivos em {tmp}")
    sys.exit(1 if total else 0)


if __name__ == "__main__":
    main()
