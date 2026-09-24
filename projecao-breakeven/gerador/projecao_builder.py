# -*- coding: utf-8 -*-
"""Motor de layout das abas de projeção (12 meses, Projetado | Realizado por mês) no padrão V4.

Cada aba visível ganha uma aba de apoio OCULTA ("_apoio_<slug>") com rótulos de mês, flags do realizado, espelho contíguo do
projetado e o funil: é dela que gráficos, KPIs e o bloco de meta leem. A aba visível fica só com o conteúdo do template.

Extensões opcionais em cfg (todas retrocompatíveis):
  acum0_key      chave da premissa "acumulado inicial" (entra no acumulado do M1)
  quarter_keys   {'receita','custo','resultado','cum'[,'vendas','roas','share']} -> bloco "Visão por quarter" (fórmulas vivas, com fonte e crescimento)
  metodologia    [(titulo, [linhas])] -> aba visível "Premissas · <aba>" (fonte, janela, alertas, premissas assumidas, cenário base)
  historico      {'meses': [{'rotulo','status','janela'}], 'linhas': [(rotulo, fmt, [valores])], 'nota': str}
  legado         mesma forma do historico, com um 4o item opcional por linha: 'sub' (linha secundaria) ou 'ating' (razao, vermelho < 1 e verde >= 1)
  envelope       {'linhas': [(rotulo, fmt, atual, mediana, melhor, melhor_mes, referencia, alvo)], 'nota': str, 'alvo_em': 'M4'}
  base_ref       {'titulo': str, 'nota': str, 'linhas': [(rotulo, fmt, [12 valores])]}   -> cenário de referência (valores do piloto)
  pilot_ref      {'header': str, 'status': str, 'valores': {indice_linha_meta: valor}}
  real_prefill   {indice_mes(0-based): {chave_metrica: valor}}
  month_labels   lista de 12 rótulos de calendário (ex.: 'set/2026')
  notes_rows     altura (linhas) do bloco de notas; padrão 8
  hidden_rows    chaves de linhas que ficam ocultas na aba (intermediárias de fórmula e flags dos cartões)
Tipos de métrica: 'input' | 'link' | 'linkf' (1º mês, demais) | 'calc' | 'calcf' (1º mês, demais) | 'cum' | 'mixed' (lista de 12: número = célula editável, texto = fórmula)
Nas fórmulas do projetado: {c} coluna do mês, {p} coluna do mês anterior, {r} coluna do realizado do mês, {HAS} flag "mês tem realizado", {P_x} premissa x.
"""
import datetime, os, re, unicodedata
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.data_source import NumFmt
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

RED, BLACK, INK, GRAY, GRAY7 = "E50914", "111111", "1F2937", "6B7280", "4B5563"
LIGHT, LINE, WHITE, GREEN = "F3F4F6", "E5E7EB", "FFFFFF", "15803D"
INPUT_FILL, LINK_FILL, INPUT_FONT = "FFF1B8", "FFFBE6", "1F2937"   # sem azul: fundo carrega o significado
REAL_FILL, REAL_LINK_FILL = "FBE0E0", "FDF2F2"                    # realizado em vermelho-claro (marca V4)
# Cor por bloco (sem azul, regra do usuário): a faixa da seção, a tarja da coluna A e a barra do rótulo saem da mesma cor,
# para quem lê a planilha separar num relance investimento, marketing (mídia) e vendas (inside sales).
BLOCOS = {
    'investimento': ("374151", "INVESTIMENTO"),
    'marketing':    ("E50914", "MARKETING"),
    'vendas':       ("B45309", "VENDAS"),
    'financeiro':   ("15803D", "FINANCEIRO"),
    'resultado':    ("111111", "RESULTADO"),
}
BLOCO_PADRAO = 'resultado'
FONT = "Calibri"
# Moeda do template. Vem do ambiente porque os formatos nascem aqui, na importação do módulo, antes de qualquer
# parsing de argumento; build_workbook lê --moeda de sys.argv e planta PROJECAO_MOEDA antes de importar este arquivo.
MOEDA = os.environ.get("PROJECAO_MOEDA", "R$")
FMT_BRL, FMT_BRL2, FMT_INT = f'"{MOEDA}" #,##0', f'"{MOEDA}" #,##0.00', '#,##0'
FMT_PCT0, FMT_PCT1, FMT_PCT2, FMT_X = '0%', '0.0%', '0.00%', '0.00"x"'

N_MONTHS_MAX = 12

def grade_colunas(n):
    """Colunas da tabela para n meses: Projetado/Realizado por mês, depois Total projetado/realizado."""
    proj = [get_column_letter(3 + 2*k) for k in range(n)]
    real = [get_column_letter(4 + 2*k) for k in range(n)]
    supc = [get_column_letter(3 + k) for k in range(n)]
    tot_p, tot_r, margem = get_column_letter(3 + 2*n), get_column_letter(4 + 2*n), get_column_letter(5 + 2*n)
    todas = ['B'] + [c for k in range(n) for c in (proj[k], real[k])] + [tot_p, tot_r]
    return proj, real, supc, tot_p, tot_r, margem, todas

def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
def fill(c): return PatternFill('solid', start_color=c, end_color=c)
from copy import copy as _copy
from openpyxl.cell.cell import MergedCell
thin = Side(style='thin', color=LINE); redbar = Side(style='medium', color=RED); whitebar = Side(style='medium', color=WHITE)
RIGHT = Alignment(horizontal='right', vertical='center'); LEFT = Alignment(horizontal='left', vertical='center', indent=1)
CENTER = Alignment(horizontal='center', vertical='center')
HWRAP = Alignment(horizontal='center', vertical='center', wrap_text=True)
HWRAPL = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
VERT = Alignment(horizontal='center', vertical='center', text_rotation=90)   # tarja do bloco na coluna A
def num(x): return f"IF(ISNUMBER({x}),{x},0)"
def slug(s): return re.sub(r"[^a-z0-9]+", "_", unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()).strip("_")

class Ctx:
    """Referências para fórmulas de KPI / meta. rng() aponta para o espelho contíguo (projetado) na aba de apoio."""
    def __init__(self, R, M, P=None, sref="", sup_last='N', tot_p='AA', proj_last='Y', n=12):
        self.R, self.M, self.P, self.sref, self.used = R, M, (P or {}), sref, []
        self.sup_last, self.tot_p, self.proj_last, self.n = sup_last, tot_p, proj_last, n
    def rng(self, key):
        self.used.append(key)
        return f"{self.sref}C{self.M[key]}:{self.sup_last}{self.M[key]}" if key in self.M else f"{self.sref}C0:{self.sup_last}0"
    def tot(self, key): return f"{self.tot_p}{self.R[key]}"
    def last(self, key): return f"{self.proj_last}{self.R[key]}"
    def prem(self, key): return self.P.get(key)

def build_sheet(ws, cfg, logo_path):
    n = int(cfg.get('n_months', N_MONTHS_MAX))
    PROJ, REAL, SUPC, TOT_P, TOT_R, MARGIN_COL, ALL_TABLE_COLS = grade_colunas(n)
    FIRST_COL, LAST_MONTH_COL = PROJ[0], REAL[-1]
    KPI_LAST = get_column_letter(max(21, 5 + 2*n))   # área dos cartões de indicadores (começam em G, 5 cartões de 3 colunas)
    wb = ws.parent
    sup = wb.create_sheet(f"_apoio_{slug(ws.title)}"[:31]); sup.sheet_state = 'hidden'
    SREF = f"'{sup.title}'!"
    def put(ref, value=None, f=None, fl=None, fmt=None, al=None, bd=None, sheet=None):
        c = (sheet or ws)[ref]
        if value is not None: c.value = value
        if f: c.font = f
        if fl: c.fill = fl
        if fmt: c.number_format = fmt
        if al: c.alignment = al
        if bd: c.border = bd
        return c
    def block_title(row, title, note=None):
        put(f'B{row}', title, font(9, True, RED))
        if note: put(f'B{row+1}', note, font(9, False, GRAY, True))

    ws.sheet_properties.tabColor = RED
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 80
    ws.column_dimensions['A'].width = 11
    ws.column_dimensions['B'].width = 52   # cabe o rótulo com prefixo de unidade sem ser cortado pela célula do Mês 1
    for col in PROJ + REAL: ws.column_dimensions[col].width = 12
    for col in ('C', 'D', 'E', 'F', 'G', 'H', 'I'): ws.column_dimensions[col].width = 15.5   # blocos de quarter/envelope legíveis
    ws.column_dimensions[TOT_P].width = 13.5; ws.column_dimensions[TOT_R].width = 13.5
    ws.column_dimensions[MARGIN_COL].width = 2
    for j in range(3, ws[f'{KPI_LAST}1'].column + 1):   # colunas usadas pelos cartões ficam legíveis
        col = get_column_letter(j)
        if (ws.column_dimensions[col].width or 0) < 12: ws.column_dimensions[col].width = 12
    sup.column_dimensions['B'].width = 44
    for col in SUPC: sup.column_dimensions[col].width = 13

    # ---- cabeçalho
    today = datetime.date.today().strftime("%d/%m/%Y")
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 26; ws.row_dimensions[3].height = 18; ws.row_dimensions[4].height = 16
    if logo_path and os.path.exists(logo_path):   # logo opcional: ponha um PNG em gerador/v4_logo.png
        img = XLImage(logo_path); img.width = 60; img.height = 61; ws.add_image(img, 'A2')
    put('B2', cfg['title'], font(18, True, BLACK), al=Alignment(vertical='center'))
    put('B3', cfg['subtitle'], font(11, False, GRAY))
    put('B4', cfg['meta_line'].format(today=today), font(9, False, GRAY, True))
    ws.row_dimensions[5].height = 3
    for col in ['A'] + ALL_TABLE_COLS: ws[f'{col}5'].fill = fill(RED)
    ws.row_dimensions[6].height = 10

    # ---- posições (aba visível)
    PREM = cfg['premises']; PREM_TOP = 9; PREM_END = PREM_TOP + len(PREM) - 1
    META_TOP = PREM_END + 2
    cur = META_TOP + 11
    QK = cfg.get('quarter_keys')
    if QK: Q_TOP = cur; cur = Q_TOP + 8
    HIST = cfg.get('historico'); ENV = cfg.get('envelope'); LEG = cfg.get('legado')
    TABLE_TITLE = cur; TABLE_NOTE = TABLE_TITLE + 1; HDR1 = TABLE_TITLE + 2; HDR2 = HDR1 + 1
    metrics = cfg['metrics']
    R = {}; r = HDR2 + 1
    for m in metrics:
        if m[0] == 'sec': r += 1; continue
        R[m[0]] = r; r += 1
    LAST_TABLE = r - 1
    P = {key: f'$C${PREM_TOP + i}' for i, (key, *_rest) in enumerate(PREM)}
    PF = {'P_' + k: v for k, v in P.items()}
    spec = {m[0]: m for m in metrics if m[0] != 'sec'}
    real_inputs = [m[0] for m in metrics if m[0] != 'sec' and m[6] == 'in']
    acum0 = P.get(cfg['acum0_key']) if cfg.get('acum0_key') else None
    charts = cfg['charts']; n_chart_rows = (len(charts) + 1) // 2
    CH_TITLE = LAST_TABLE + 2; CH_TOP = CH_TITLE + 2; ROW_STEP = 19
    LAST_ROW = CH_TOP + ROW_STEP * n_chart_rows

    # ---- posições (aba de apoio)
    S_MONTH, S_HAS, S_LAST, S_POS, S_NEG = 2, 3, 4, 5, 6
    probe = Ctx(R, {}, P, SREF, SUPC[-1], TOT_P, PROJ[-1], n); cfg['kpis'](probe); cfg['meta'](probe, ['X1'] * 8)
    mirror_keys = []
    chart_keys = [s[0] for sp in charts if sp['type'] == 'line' for s in sp['series']] + [sp[w] for sp in charts if sp['type'] == 'growth' for w in ('bars', 'line')]
    for k in [cfg['cum_key'], 'flag'] + chart_keys + probe.used + list(cfg.get('espelho') or []):   # espelho: linhas extras para gráficos de fora da aba
        if k and k in R and k not in mirror_keys: mirror_keys.append(k)
    M_TOP = S_NEG + 2
    M = {k: M_TOP + i for i, k in enumerate(mirror_keys)}
    F0 = M_TOP + len(mirror_keys) + 2; FUNNEL = cfg['funnel']
    ctx = Ctx(R, M, P, SREF, SUPC[-1], TOT_P, PROJ[-1], n)
    def has(k): return f"{SREF}${SUPC[k]}${S_HAS}"

    # ---- 01 premissas gerais
    put('B7', "01 · PREMISSAS GERAIS", font(9, True, RED)); put('E7', "02 · RESULTADO EXECUTIVO (PROJETADO)", font(9, True, RED))
    put('B8', "Células amarelas são editáveis; as brancas são calculadas pelo modelo.", font(9, False, GRAY, True))
    put('E8', "Indicadores calculados a partir da projeção da tabela mês a mês (seção 03).", font(9, False, GRAY, True))
    for i, (key, label, val, fmt) in enumerate(PREM):
        rr = PREM_TOP + i
        ws.row_dimensions[rr].height = 22 if rr in (10, 14) else 18
        put(f'B{rr}', label, font(10, False, INK), al=LEFT, bd=Border(bottom=thin, left=redbar))
        put(f'C{rr}', val, font(10, True, INPUT_FONT), fill(INPUT_FILL), fmt, RIGHT, Border(bottom=thin))
    if cfg.get('budget_day'):
        mk, dk = cfg['budget_day']
        put(f'B{PREM_END+1}', "Budget médio de mídia por dia (Mês 1)", font(9, False, GRAY, True), al=LEFT)
        put(f'C{PREM_END+1}', f"=IFERROR({P[mk]}/{P[dk]},0)", font(9, False, GRAY, True), fmt=FMT_BRL2, al=RIGHT)

    # ---- 02 KPI cards (3 colunas cada, a partir de G: deixa respiro entre as premissas e os cartões)
    KPI_COL0, cards = 7, cfg['kpis'](ctx)
    POR_LINHA = 5 if len(cards) > 8 else 4   # no máximo duas fileiras; 7 cartões ficam 4 + 3, não 5 + 2
    for i, (label, formula, fmt, sub) in enumerate(cards):
        col0 = KPI_COL0 + 3 * (i % POR_LINHA)
        c1, c2 = get_column_letter(col0), get_column_letter(col0 + 2)
        r0 = 9 + 4 * (i // POR_LINHA)
        cols = [get_column_letter(j) for j in range(ws[c1 + '1'].column, ws[c2 + '1'].column + 1)]
        for rr in (r0, r0 + 1, r0 + 2):
            ws.merge_cells(f'{c1}{rr}:{c2}{rr}')
            for cc in cols:
                ws[f'{cc}{rr}'].fill = fill(LIGHT)
                ws[f'{cc}{rr}'].border = Border(left=whitebar if cc == c1 else None)
        put(f'{c1}{r0}', label, font(8, True, GRAY), al=LEFT)
        v = put(f'{c1}{r0+1}', formula, font(14, True, BLACK), al=LEFT)
        if fmt: v.number_format = fmt
        put(f'{c1}{r0+2}', sub, font(8, False, GRAY, True), al=LEFT)
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    put('G17', "Legenda", font(8, True, GRAY), al=LEFT)
    for sw, tx, fl, text in [('H', 'I', INPUT_FILL, "Alavanca editável (amarelo)"), ('K', 'L', LINK_FILL, "Puxa a premissa geral"),
                             ('N', 'O', REAL_FILL, "Realizado: preencha (vermelho-claro)"), ('Q', 'R', REAL_LINK_FILL, "Realizado: puxa o projetado"),
                             ('T', 'U', None, "Calculado pelo modelo")]:
        put(f'{sw}17', None, None, fill(fl) if fl else None, bd=box); put(f'{tx}17', text, font(8, False, GRAY), al=LEFT)

    # ---- meta de breakeven + coluna PILOTO + notas
    put(f'B{META_TOP}', "META DE BREAKEVEN · ENGENHARIA REVERSA (PROJETADO)", font(9, True, RED))
    mcells = [f'C{META_TOP + 1 + i}' for i in range(8)]
    for i, (label, val, fmt, is_input) in enumerate(cfg['meta'](ctx, mcells)):
        rr = META_TOP + 1 + i; ws.row_dimensions[rr].height = 20
        put(f'B{rr}', label, font(10, False, INK), al=LEFT, bd=Border(bottom=thin, left=redbar if is_input else None))
        if is_input: put(f'C{rr}', val, font(10, True, INPUT_FONT), fill(INPUT_FILL), fmt, RIGHT, Border(bottom=thin))
        else: put(f'C{rr}', val, font(10, True, INK), None, fmt, RIGHT, Border(bottom=thin))
    gap = mcells[-1]
    ws.conditional_formatting.add(gap, CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=FONT, bold=True, color=RED)))
    ws.conditional_formatting.add(gap, CellIsRule(operator='lessThanOrEqual', formula=['0'], font=Font(name=FONT, bold=True, color=GREEN)))
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2=str(n), showErrorMessage=True, errorTitle='Mês inválido', error=f'Informe um mês de 1 a {n}.')
    ws.add_data_validation(dv); dv.add(mcells[0])
    PR = cfg.get('pilot_ref')
    if PR:
        put(f'D{META_TOP}', PR.get('header', "PILOTO"), font(8, True, GRAY), al=RIGHT)
        st = PR.get('status') or ""
        put(f'D{META_TOP+1}', st, font(9, True, GREEN if st.startswith('REALISTA') else RED), al=RIGHT, bd=Border(bottom=thin))
        for idx, val in PR.get('valores', {}).items():
            put(f'D{META_TOP+1+idx}', val, font(9, False, GRAY, True), fmt=FMT_INT if isinstance(val, (int, float)) else None, al=RIGHT, bd=Border(bottom=thin))

    # ---- visão por quarter (fórmulas vivas sobre a tabela)
    if QK:
        block_title(Q_TOP, "VISÃO POR QUARTER (PROJETADO)", "Soma dos três meses de cada quarter; o acumulado é a posição no fim do quarter. A venda pode cair num mês e o trimestre absorve.")
        heads = ["QUARTER", "MESES", "FONTE", "RECEITA", "CRESC. VS QUARTER ANTERIOR", "CUSTO (FEE + MÍDIA)", "RESULTADO LÍQUIDO", "ACUMULADO NO FIM"]
        qcols = ['B', 'C', 'D', 'E', 'F', 'G', 'H']
        ws.row_dimensions[Q_TOP+2].height = 34
        for c, h in zip(qcols, heads): put(f'{c}{Q_TOP+2}', h, font(8, True, WHITE), fill(BLACK), al=HWRAP if c != 'B' else HWRAPL)
        for q in range(4):
            rr = Q_TOP + 3 + q; cols3 = PROJ[3*q:3*q+3]; flags3 = [f"{SREF}{SUPC[3*q+i]}${S_HAS}" for i in range(3)]
            put(f'B{rr}', f"Q{q+1}", font(9, True, INK), al=LEFT, bd=Border(bottom=thin))
            put(f'C{rr}', f'=INDEX({SREF}$C${S_MONTH}:$N${S_MONTH},1,{3*q+1})&" a "&INDEX({SREF}$C${S_MONTH}:$N${S_MONTH},1,{3*q+3})', font(9, False, GRAY), None, None, CENTER, Border(bottom=thin))
            n_real = "+".join(f"IF({f}>0,1,0)" for f in flags3)
            put(f'D{rr}', f'=IF({n_real}=3,"realizado",IF({n_real}=0,"projetado","misto"))', font(9, False, GRAY), None, None, CENTER, Border(bottom=thin))
            for c, key, fmt in [('E', QK['receita'], FMT_BRL), ('G', QK['custo'], FMT_BRL), ('H', QK['resultado'], FMT_BRL)]:
                put(f'{c}{rr}', "=" + "+".join(f"{cc}{R[key]}" for cc in cols3), font(9, key == QK['resultado'], INK), None, fmt, RIGHT, Border(bottom=thin))
            put(f'F{rr}', "" if q == 0 else f'=IFERROR(E{rr}/E{rr-1}-1,"")', font(9, False, INK), None, FMT_PCT1, RIGHT, Border(bottom=thin))
            put(f'I{rr}', f"={cols3[-1]}{R[QK['cum']]}", font(9, True, INK), None, FMT_BRL, RIGHT, Border(bottom=thin))
        put(f'I{Q_TOP+2}', "ACUMULADO NO FIM", font(8, True, WHITE), fill(BLACK), al=HWRAP)
        for c in ('F', 'H', 'I'):
            rng = f"{c}{Q_TOP+3}:{c}{Q_TOP+6}"
            ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'], font=Font(name=FONT, size=9, bold=True, color=RED)))
            ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=['0'], font=Font(name=FONT, size=9, bold=True, color=GREEN)))

    # ---- 03 tabela
    put(f'B{TABLE_TITLE}', "03 · PROJEÇÃO MÊS A MÊS · PROJETADO × REALIZADO", font(9, True, RED))
    put(f'B{TABLE_NOTE}', "Cada mês tem duas colunas: Projetado (modelo) e Realizado (preencha volumes e valores nas células vermelho-claro; taxas, custos unitários, acumulados e resultado do realizado saem por fórmula e ficam em branco até o mês ter dado). A coluna Total soma volumes/valores e mostra a taxa efetiva para percentuais.", font(9, False, GRAY, True))
    ws.row_dimensions[HDR1].height = 20; ws.row_dimensions[HDR2].height = 18
    ws.merge_cells(f'B{HDR1}:B{HDR2}'); put(f'B{HDR1}', "MÉTRICA", font(9, True, WHITE), fill(BLACK), al=LEFT); ws[f'B{HDR2}'].fill = fill(BLACK)
    labels = cfg.get('month_labels')
    for k in range(n):
        ws.merge_cells(f'{PROJ[k]}{HDR1}:{REAL[k]}{HDR1}')
        head = f"MÊS {k+1}" + (f" · {labels[k]}" if labels else "")
        put(f'{PROJ[k]}{HDR1}', head, font(9, True, WHITE), fill(BLACK), al=CENTER); ws[f'{REAL[k]}{HDR1}'].fill = fill(BLACK)
        put(f'{PROJ[k]}{HDR2}', "Projetado", font(8, True, WHITE), fill(BLACK), al=CENTER)
        put(f'{REAL[k]}{HDR2}', "Realizado", font(8, True, WHITE), fill(GRAY7), al=CENTER)
    ws.merge_cells(f'{TOT_P}{HDR1}:{TOT_R}{HDR1}')
    put(f'{TOT_P}{HDR1}', "TOTAL / MÉDIA", font(9, True, WHITE), fill(BLACK), al=CENTER); ws[f'{TOT_R}{HDR1}'].fill = fill(BLACK)
    put(f'{TOT_P}{HDR2}', "Projetado", font(8, True, WHITE), fill(BLACK), al=CENTER)
    put(f'{TOT_R}{HDR2}', "Realizado", font(8, True, WHITE), fill(GRAY7), al=CENTER)

    def lst(cols, rr): return ",".join(f"{c}{rr}" for c in cols)
    def total_proj(spec_t, key):
        rr = R[key]
        if spec_t == 'sum': return f"=SUM({lst(PROJ, rr)})"
        if spec_t == 'avg': return f"=AVERAGE({lst(PROJ, rr)})"
        if spec_t == 'last': return f"={PROJ[-1]}{rr}"
        if spec_t.startswith('div1000:'):
            a, b = spec_t[8:].split('/'); return f"=IFERROR({TOT_P}{R[a]}/{TOT_P}{R[b]}*1000,0)"
        if spec_t.startswith('div:'):
            a, b = spec_t[4:].split('/'); return f"=IFERROR({TOT_P}{R[a]}/{TOT_P}{R[b]},0)"
    def total_real(spec_t, key):
        rr = R[key]; L = lst(REAL, rr)
        if spec_t == 'sum': return f'=IF(COUNT({L})=0,"",SUM({L}))'
        if spec_t == 'avg': return f'=IFERROR(AVERAGE({L}),"")'
        if spec_t == 'last': return f'=IF({SREF}$C${S_LAST}=0,"",INDEX({FIRST_COL}{rr}:{LAST_MONTH_COL}{rr},1,2*{SREF}$C${S_LAST}))'
        if spec_t.startswith('div1000:'):
            a, b = spec_t[8:].split('/'); return f'=IFERROR({TOT_R}{R[a]}/{TOT_R}{R[b]}*1000,"")'
        if spec_t.startswith('div:'):
            a, b = spec_t[4:].split('/'); return f'=IFERROR({TOT_R}{R[a]}/{TOT_R}{R[b]},"")'

    bold_rows = {k for k in cfg.get('bold_rows', ()) if k in R}; r = HDR2 + 1
    bloco, bloco_ini, tarjas = BLOCO_PADRAO, None, []
    for m in metrics:
        ws.row_dimensions[r].height = 17
        if m[0] == 'sec':
            novo = (m[2] if len(m) > 2 else None) or BLOCO_PADRAO
            if bloco_ini is not None and novo != bloco:   # seções vizinhas do mesmo bloco continuam na mesma tarja
                tarjas.append((bloco_ini, r - 1, bloco)); bloco_ini = None
            bloco = novo
            cor_sec = BLOCOS.get(bloco, BLOCOS[BLOCO_PADRAO])[0]
            for col in ALL_TABLE_COLS:
                ws[f'{col}{r}'].fill = fill(cor_sec); ws[f'{col}{r}'].border = Border(bottom=thin)
            put(f'B{r}', m[1], font(8, True, WHITE), al=LEFT)
            if bloco_ini is None: bloco_ini = r
            r += 1; continue
        key, label, kind, src, fmt, tot, real = m
        is_bold = key in bold_rows
        is_cum0 = (kind == 'cum' and key == cfg['cum_key'] and acum0 is not None)
        cor_bloco = BLOCOS.get(bloco, BLOCOS[BLOCO_PADRAO])[0]
        # alavanca: rótulo e meses projetados na cor da alavanca, com barra da cor do bloco (o total é fórmula e fica cinza)
        mix_edit = kind == 'mixed' and any(not (isinstance(x, str) and x.startswith('=')) for x in src)   # algum mês digitável
        mix_link = kind == 'mixed' and any(isinstance(x, str) and x.startswith('={P_') for x in src)        # algum mês puxa a premissa
        alav = (fill(INPUT_FILL) if kind == 'input' or mix_edit else
                fill(LINK_FILL) if kind in ('link', 'linkf') or mix_link else None)
        put(f'B{r}', label, font(9, True, INK), alav, None, LEFT,
            Border(bottom=thin, left=Side(style='medium', color=cor_bloco) if alav is not None else None))
        for k in range(n):
            pc, rc = PROJ[k], REAL[k]
            pp, rp = (PROJ[k-1], REAL[k-1]) if k > 0 else (None, None)
            FMTK = dict(c=pc, p=pp, r=rc, HAS=has(k), L=PROJ[-1], **R, **PF)
            if kind == 'mixed':
                v = src[k]
                if isinstance(v, str) and v.startswith('='):
                    puxa = v.startswith('={P_')
                    put(f'{pc}{r}', v.format(**FMTK), font(9, is_bold and not puxa, INPUT_FONT if puxa else INK, puxa),
                        fill(LINK_FILL) if puxa else None, fmt, RIGHT, Border(bottom=thin))
                else:
                    put(f'{pc}{r}', v, font(9, True, INPUT_FONT), fill(INPUT_FILL), fmt, RIGHT, Border(bottom=thin))
            elif kind == 'input':
                put(f'{pc}{r}', src[k], font(9, True, INPUT_FONT), fill(INPUT_FILL), fmt, RIGHT, Border(bottom=thin))
            elif kind == 'link':
                put(f'{pc}{r}', f"={P[src]}", font(9, False, INPUT_FONT, True), fill(LINK_FILL), fmt, RIGHT, Border(bottom=thin))
            elif kind == 'linkf':
                put(f'{pc}{r}', (src[0] if k == 0 else src[1]).format(**FMTK), font(9, False, INPUT_FONT, True), fill(LINK_FILL), fmt, RIGHT, Border(bottom=thin))
            elif kind == 'cum':
                if k == 0: f = f"={acum0}+{pc}{R[src]}" if is_cum0 else f"={pc}{R[src]}"
                else: f = f"={pp}{r}+{pc}{R[src]}"
                put(f'{pc}{r}', f, font(9, is_bold, INK), None, fmt, RIGHT, Border(bottom=thin))
            elif kind == 'calcf':
                put(f'{pc}{r}', (src[0] if k == 0 else src[1]).format(**FMTK), font(9, is_bold, INK), None, fmt, RIGHT, Border(bottom=thin))
            else:
                put(f'{pc}{r}', src.format(**FMTK), font(9, is_bold, INK), None, fmt, RIGHT, Border(bottom=thin))
            if real == '""':
                # linha calculada: não há o que digitar, então o par Projetado|Realizado
                # vira uma célula só. O valor já é o realizado nos meses que têm dado.
                cp = ws[f'{pc}{r}']
                put(f'{rc}{r}', None, _copy(cp.font), _copy(cp.fill), fmt, CENTER, Border(bottom=thin))
                ws.merge_cells(f'{pc}{r}:{rc}{r}')
                cp.alignment = CENTER
            elif real == 'in':
                put(f'{rc}{r}', None, font(9, True, INPUT_FONT), fill(REAL_FILL), fmt, RIGHT, Border(bottom=thin))
            elif real == 'proj':
                put(f'{rc}{r}', f'=IF({has(k)}=0,"",{pc}{r})', font(9, False, INPUT_FONT, True), fill(REAL_LINK_FILL), fmt, RIGHT, Border(bottom=thin))
            elif kind == 'cum' and real == 'cum':
                base = f"{rc}{R[src]}"
                if k == 0: expr = (f"{num(acum0)}+" if is_cum0 else "") + num(base)
                else: expr = f"{num(f'{rp}{r}')}+{num(base)}"
                put(f'{rc}{r}', f'=IF({has(k)}=0,"",IFERROR({expr},""))', font(9, is_bold, INK), None, fmt, RIGHT, Border(bottom=thin))
            else:
                expr = real.format(c=rc, p=rp, pc=pc, **R, **PF)
                put(f'{rc}{r}', f'=IF({has(k)}=0,"",IFERROR({expr},""))', font(9, is_bold, INK), None, fmt, RIGHT, Border(bottom=thin))
        al_tot = CENTER if real == '""' else RIGHT
        fl_tot = fill(LIGHT)   # o total é sempre fórmula: fica cinza, porque amarelo é o código de "você pode editar"
        put(f'{TOT_P}{r}', total_proj(tot, key), font(9, True, INK), fl_tot, fmt, al_tot, Border(bottom=thin))
        if real == '""':
            put(f'{TOT_R}{r}', None, font(9, True, INK), fl_tot, fmt, al_tot, Border(bottom=thin))
            ws.merge_cells(f'{TOT_P}{r}:{TOT_R}{r}')
        else:
            put(f'{TOT_R}{r}', total_real(tot, key), font(9, True, INK), fl_tot, fmt, RIGHT, Border(bottom=thin))
        r += 1
    for key in [k for k in cfg.get('hidden_rows', ()) if k in R]:
        ws.row_dimensions[R[key]].hidden = True   # intermediária de fórmula: existe para a planilha, não para quem lê
    if bloco_ini is not None: tarjas.append((bloco_ini, r - 1, bloco))
    for ini, fim, b in tarjas:   # tarja vertical na coluna A: o bloco da linha se lê sem precisar procurar o título
        cor, nome = BLOCOS.get(b, BLOCOS[BLOCO_PADRAO])
        for _rr in range(ini, fim + 1): ws[f'A{_rr}'].fill = fill(cor)
        put(f'A{ini}', nome, font(8, True, WHITE), fill(cor), None, VERT)
        if fim > ini: ws.merge_cells(f'A{ini}:A{fim}')
    for k, vals in (cfg.get('real_prefill') or {}).items():
        for key, v in vals.items():
            cel = ws[f'{REAL[k]}{R[key]}'] if key in R else None
            if cel is not None and v is not None and not isinstance(cel, MergedCell): cel.value = v
    for key in [k for k in cfg['signed_rows'] if k in R]:
        rng = f"{FIRST_COL}{R[key]}:{TOT_R}{R[key]}"
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'], font=Font(name=FONT, size=9, bold=True, color=RED)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=['0'], font=Font(name=FONT, size=9, bold=True, color=GREEN)))
    for key in [k for k in cfg.get('x_rows', ()) if k in R]:   # múltiplo: o prejuízo começa abaixo de 1,0x, não abaixo de zero
        rng = f"{FIRST_COL}{R[key]}:{TOT_R}{R[key]}"
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['1'], font=Font(name=FONT, size=9, bold=True, color=RED)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=['1'], font=Font(name=FONT, size=9, bold=True, color=GREEN)))

    # ---- aba de apoio (oculta): rótulos, flags do realizado, espelho contíguo, funil
    put('B1', f"Apoio de '{ws.title}' (gerado; gráficos, KPIs, meta e realizado dependem desta aba)", font(9, True, RED), sheet=sup)
    put(f'B{S_MONTH}', "Rótulo do mês", font(8, False, GRAY), al=LEFT, sheet=sup)
    put(f'B{S_HAS}', "Realizado preenchido? (nº do mês ou 0)", font(8, False, GRAY), al=LEFT, sheet=sup)
    put(f'B{S_LAST}', "Último mês com realizado", font(8, False, GRAY), al=LEFT, sheet=sup)
    put(f'B{S_POS}', "Acumulado projetado positivo", font(8, False, GRAY), al=LEFT, sheet=sup)
    put(f'B{S_NEG}', "Acumulado projetado negativo", font(8, False, GRAY), al=LEFT, sheet=sup)
    cum_key = cfg['cum_key']; MAIN = f"'{ws.title}'!"
    for k, col in enumerate(SUPC):
        put(f'{col}{S_MONTH}', f"M{k+1}", font(8, False, GRAY), al=CENTER, sheet=sup)
        put(f'{col}{S_HAS}', f"=IF(COUNT({','.join(f'{MAIN}{REAL[k]}{R[x]}' for x in real_inputs)})>0,{k+1},0)", font(8, False, GRAY), al=CENTER, sheet=sup)
        put(f'{col}{S_POS}', f"=MAX({MAIN}{PROJ[k]}{R[cum_key]},0)", font(8, False, GRAY), fmt=FMT_BRL, al=RIGHT, sheet=sup)
        put(f'{col}{S_NEG}', f"=MIN({MAIN}{PROJ[k]}{R[cum_key]},0)", font(8, False, GRAY), fmt=FMT_BRL, al=RIGHT, sheet=sup)
    put(f'C{S_LAST}', f"=MAX(C{S_HAS}:{SUPC[-1]}{S_HAS})", font(8, False, GRAY), al=CENTER, sheet=sup)
    put(f'B{M_TOP-1}', "Espelho contíguo do projetado", font(8, True, GRAY), al=LEFT, sheet=sup)
    for key, mr in M.items():
        put(f'B{mr}', spec[key][1], font(8, False, GRAY), al=LEFT, sheet=sup)
        for k, col in enumerate(SUPC):
            put(f'{col}{mr}', f"={MAIN}{PROJ[k]}{R[key]}", font(8, False, GRAY), fmt=spec[key][4], al=RIGHT, sheet=sup)
    PBX = cfg.get('payback_estendido')
    PB0 = F0 + len(cfg['funnel']) + 2
    if PBX:   # a tabela tem 12 colunas, mas a curva de payback pode precisar de mais meses para cruzar o zero
        put(f'B{PB0-1}', "Curva de payback estendida (tabela + meses à frente, até o acumulado zerar)", font(8, True, GRAY), al=LEFT, sheet=sup)
        put(f'B{PB0}', "Mês", font(8, False, GRAY), al=LEFT, sheet=sup)
        put(f'B{PB0+1}', "Acumulado positivo", font(8, False, GRAY), al=LEFT, sheet=sup)
        put(f'B{PB0+2}', "Acumulado negativo", font(8, False, GRAY), al=LEFT, sheet=sup)
        for k, col in enumerate(SUPC):   # os 12 meses da tabela vêm por fórmula, para acompanhar quem editar a aba
            put(f'{col}{PB0}', (cfg.get('month_labels') or [f"M{k+1}"] * len(SUPC))[k], font(8, False, GRAY), al=CENTER, sheet=sup)
            put(f'{col}{PB0+1}', f"=MAX({MAIN}{PROJ[k]}{R[cum_key]},0)", font(8, False, GRAY), fmt=FMT_BRL, al=RIGHT, sheet=sup)
            put(f'{col}{PB0+2}', f"=MIN({MAIN}{PROJ[k]}{R[cum_key]},0)", font(8, False, GRAY), fmt=FMT_BRL, al=RIGHT, sheet=sup)
        for j, (rot, ac) in enumerate(zip(PBX['rotulos'], PBX['acumulado'])):
            col = get_column_letter(3 + len(SUPC) + j)
            put(f'{col}{PB0}', rot, font(8, False, GRAY), al=CENTER, sheet=sup)
            put(f'{col}{PB0+1}', max(ac, 0), font(8, False, GRAY), fmt=FMT_BRL, al=RIGHT, sheet=sup)
            put(f'{col}{PB0+2}', min(ac, 0), font(8, False, GRAY), fmt=FMT_BRL, al=RIGHT, sheet=sup)
    put(f'B{F0-1}', "Funil acumulado projetado · 12 meses", font(8, True, GRAY), al=LEFT, sheet=sup)
    put(f'C{F0-1}', "Volume", font(8, True, GRAY), al=RIGHT, sheet=sup); put(f'D{F0-1}', "Conv. etapa", font(8, True, GRAY), al=RIGHT, sheet=sup)
    for i, (lab, key) in enumerate(FUNNEL):
        rr = F0 + i
        put(f'B{rr}', lab, font(8, False, GRAY), al=LEFT, sheet=sup)
        put(f'C{rr}', f"={MAIN}{TOT_P}{R[key]}", font(8, False, GRAY), fmt=FMT_INT, al=RIGHT, sheet=sup)
        put(f'D{rr}', "" if i == 0 else f"=IFERROR(C{rr}/C{rr-1},0)", font(8, False, GRAY), fmt=FMT_PCT1, al=RIGHT, sheet=sup)

    # ---- 04 gráficos (leem a aba de apoio)
    put(f'B{CH_TITLE}', "04 · GRÁFICOS", font(9, True, RED))
    put(f'B{CH_TITLE+1}', "Leem a projeção da seção 03 e acompanham qualquer alteração nas premissas.", font(9, False, GRAY, True))
    cats = Reference(sup, min_col=3, max_col=2 + n, min_row=S_MONTH, max_row=S_MONTH)
    CW, CHH = 18.5, 8.6
    def grid(ax):
        ax.delete = False
        if ax.majorGridlines is not None:
            ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    def finish(ch):
        ch.visible_cells_only = False
        return ch
    def line_chart(sp):
        ch = LineChart(); ch.title = sp['title']; ch.height = CHH; ch.width = CW; ch.legend.position = 'b'
        for key, color, dash in sp['series']:
            ch.add_data(Reference(sup, min_col=2, max_col=2 + n, min_row=M[key], max_row=M[key]), from_rows=True, titles_from_data=True)
        for s, (key, color, dash) in zip(ch.series, sp['series']):
            s.smooth = False
            s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = 22000
            if dash: s.graphicalProperties.line.dashStyle = 'sysDash'
            s.marker.symbol = 'circle'; s.marker.size = 5
            s.marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color))
        ch.set_categories(cats); grid(ch.x_axis); grid(ch.y_axis)
        ch.y_axis.numFmt = NumFmt(formatCode=sp['y_fmt'], sourceLinked=False)
        return finish(ch)
    def payback_chart(sp):
        ch = BarChart(); ch.type = 'col'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 60
        ch.title = sp['title']; ch.height = CHH; ch.width = CW; ch.legend = None
        if PBX:   # curva até o payback, mesmo quando ele cai depois das 12 colunas da tabela
            nx = n + len(PBX['rotulos'])
            ch.add_data(Reference(sup, min_col=2, max_col=2 + nx, min_row=PB0 + 1, max_row=PB0 + 1), from_rows=True, titles_from_data=True)
            ch.add_data(Reference(sup, min_col=2, max_col=2 + nx, min_row=PB0 + 2, max_row=PB0 + 2), from_rows=True, titles_from_data=True)
            cats_pb = Reference(sup, min_col=3, max_col=2 + nx, min_row=PB0, max_row=PB0)
        else:
            ch.add_data(Reference(sup, min_col=2, max_col=2 + n, min_row=S_POS, max_row=S_POS), from_rows=True, titles_from_data=True)
            ch.add_data(Reference(sup, min_col=2, max_col=2 + n, min_row=S_NEG, max_row=S_NEG), from_rows=True, titles_from_data=True)
            cats_pb = cats
        for s, color in zip(ch.series, (GREEN, RED)):
            s.graphicalProperties.solidFill = color; s.graphicalProperties.line.solidFill = color
        ch.set_categories(cats_pb); grid(ch.x_axis); grid(ch.y_axis)
        ch.y_axis.numFmt = NumFmt(formatCode=FMT_BRL, sourceLinked=False); ch.x_axis.tickLblPos = 'low'
        return finish(ch)
    def funnel_chart(sp):
        n = len(FUNNEL)
        ch = BarChart(); ch.type = 'bar'; ch.title = sp['title']; ch.height = CHH; ch.width = CW; ch.legend = None; ch.gapWidth = 45
        ch.add_data(Reference(sup, min_col=3, max_col=3, min_row=F0, max_row=F0 + n - 1), titles_from_data=False)
        ch.set_categories(Reference(sup, min_col=2, max_col=2, min_row=F0, max_row=F0 + n - 1))
        ch.series[0].graphicalProperties.solidFill = RED; ch.series[0].graphicalProperties.line.solidFill = RED
        ch.x_axis.scaling.orientation = 'maxMin'; grid(ch.x_axis); grid(ch.y_axis)
        ch.y_axis.numFmt = NumFmt(formatCode=FMT_INT, sourceLinked=False)
        ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.numFmt = FMT_INT
        return finish(ch)
    def growth_chart(sp):
        ch = BarChart(); ch.type = 'col'; ch.title = sp['title']; ch.height = CHH; ch.width = CW; ch.gapWidth = 40
        ch.add_data(Reference(sup, min_col=2, max_col=2 + n, min_row=M[sp['bars']], max_row=M[sp['bars']]), from_rows=True, titles_from_data=True)
        ch.series[0].graphicalProperties.solidFill = RED; ch.series[0].graphicalProperties.line.solidFill = RED
        ch.set_categories(cats); grid(ch.x_axis); grid(ch.y_axis)
        ch.y_axis.numFmt = NumFmt(formatCode=FMT_BRL, sourceLinked=False)
        ln = LineChart()
        ln.add_data(Reference(sup, min_col=2, max_col=2 + n, min_row=M[sp['line']], max_row=M[sp['line']]), from_rows=True, titles_from_data=True)
        ln.series[0].graphicalProperties.line.solidFill = BLACK; ln.series[0].graphicalProperties.line.width = 20000
        ln.series[0].smooth = False; ln.series[0].marker.symbol = 'circle'; ln.series[0].marker.size = 5
        ln.y_axis.axId = 200; ln.y_axis.title = None; ln.y_axis.numFmt = NumFmt(formatCode=FMT_PCT0, sourceLinked=False)
        ln.y_axis.majorGridlines = None; ln.y_axis.crosses = 'max'
        ch.y_axis.crosses = 'autoZero'; ch += ln
        ch.legend.position = 'b'
        return finish(ch)
    makers = {'line': line_chart, 'payback': payback_chart, 'funnel': funnel_chart, 'growth': growth_chart}
    for i, sp in enumerate(charts):
        anchor_col = 'B' if i % 2 == 0 else 'J'; anchor_row = CH_TOP + ROW_STEP * (i // 2)
        ws.add_chart(makers[sp['type']](sp), f'{anchor_col}{anchor_row}')

    # ---- aba de metodologia (fonte, janela, alertas, premissas assumidas, cenário base)
    MET = cfg.get('metodologia'); BASE = cfg.get('base_ref')
    if MET or BASE or HIST or ENV or LEG:
        met = wb.create_sheet(f"Premissas · {ws.title}"[:31])
        met.sheet_view.showGridLines = False
        met.column_dimensions['A'].width = 3; met.column_dimensions['B'].width = 46
        for j in range(3, 3 + max(len(SUPC), 7) + 18): met.column_dimensions[get_column_letter(j)].width = 14
        met.row_dimensions[1].height = 8
        put('B2', f"Premissas · {cfg['title']}", font(16, True, BLACK), al=Alignment(vertical='center'), sheet=met)
        put('B3', cfg['subtitle'], font(10, False, GRAY), sheet=met)
        met.row_dimensions[4].height = 3
        for col in ['B'] + SUPC: met[f'{col}4'].fill = fill(RED)
        rr = 6
        WRAP = Alignment(wrap_text=True, vertical='top', horizontal='left', indent=1)
        def secoes(lista, rr):
          for titulo, linhas in (lista or []):
              put(f'B{rr}', titulo, font(9, True, RED), sheet=met); rr += 1
              if isinstance(linhas, dict):   # tabela: {'colunas': [...], 'linhas': [[...]], 'spans': [(col_ini, col_fim), ...], 'nota': str}
                  spans = linhas.get('spans') or [('B', 'B'), ('C', 'F'), ('G', 'H'), ('I', 'J'), ('K', 'N')][:len(linhas['colunas'])]
                  def larg(a, b):  # caracteres que cabem no bloco mesclado
                      return sum(met.column_dimensions[get_column_letter(j)].width for j in range(met[a + '1'].column, met[b + '1'].column + 1)) * 1.1
                  if linhas.get('nota'):
                      met.merge_cells(f'B{rr}:N{rr}'); put(f'B{rr}', linhas['nota'], font(9, False, GRAY, True), al=WRAP, sheet=met)
                      met.row_dimensions[rr].height = max(16, 13 * (1 + len(linhas['nota']) // 150)); rr += 1
                  for j, (a, b) in enumerate(spans):
                      if a != b: met.merge_cells(f'{a}{rr}:{b}{rr}')
                      put(f'{a}{rr}', linhas['colunas'][j], font(8, True, WHITE), fill(BLACK), al=HWRAPL, sheet=met)
                      for jj in range(met[a + '1'].column + 1, met[b + '1'].column + 1): met[f'{get_column_letter(jj)}{rr}'].fill = fill(BLACK)
                  met.row_dimensions[rr].height = 22; rr += 1
                  for lin in linhas['linhas']:
                      altura = 1
                      for j, (a, b) in enumerate(spans):
                          txt = '' if lin[j] is None else str(lin[j])
                          if a != b: met.merge_cells(f'{a}{rr}:{b}{rr}')
                          put(f'{a}{rr}', txt, font(9, j == 0, INK), al=WRAP, bd=Border(bottom=thin), sheet=met)
                          for jj in range(met[a + '1'].column + 1, met[b + '1'].column + 1): met[f'{get_column_letter(jj)}{rr}'].border = Border(bottom=thin)
                          linhas_txt = sum(1 + len(par) // max(1, int(larg(a, b))) for par in txt.split('\n'))
                          altura = max(altura, linhas_txt)
                      met.row_dimensions[rr].height = max(16, 12.5 * altura + 4); rr += 1
                  rr += 1; continue
              for ln in linhas:
                  link = None
                  if isinstance(ln, dict): ln, link = ln.get('texto', ''), ln.get('link')   # linha com hyperlink (fontes)
                  met.merge_cells(f'B{rr}:N{rr}')
                  met.row_dimensions[rr].height = max(16, 13 * (1 + len(ln) // 150))
                  c_ = put(f'B{rr}', ln, font(9, False, INK), al=WRAP, sheet=met)
                  if link:
                      c_.hyperlink = link; c_.font = Font(name=FONT, size=9, color=INK, underline='single')
                  rr += 1
              rr += 1
          return rr
        rr = secoes(MET, rr)
        def tabela_meses(titulo, nota, meses_, linhas_, rr):
            """Tabela mês a mês da aba Premissas. Serve ao histórico e ao legado (projetado × realizado)."""
            put(f'B{rr}', titulo, font(9, True, RED), sheet=met)
            if nota: put(f'B{rr+1}', nota, font(9, False, GRAY, True), sheet=met)
            hh = rr + 2; met.row_dimensions[hh].height = 20
            hcols = [get_column_letter(3 + j) for j in range(len(meses_))]
            put(f'B{hh}', "MÉTRICA", font(9, True, WHITE), fill(BLACK), al=HWRAPL, sheet=met)
            for j, mth in enumerate(meses_):
                put(f'{hcols[j]}{hh}', mth['rotulo'], font(8, True, WHITE), fill(BLACK), al=HWRAP, sheet=met)
            tem_status = any(m.get('status') or m.get('janela') or m.get('referencia') or m.get('pre') for m in meses_)
            st = hh   # sem status (legado), as linhas começam logo abaixo do cabeçalho
            if tem_status:
                st = hh + 1; met.row_dimensions[st].height = 28
                put(f'B{st}', "Status na fonte", font(9, True, INK), al=LEFT, bd=Border(bottom=thin), sheet=met)
                for j, mth in enumerate(meses_):
                    tags = [t for t, on in (("janela", mth.get('janela')), ("referência", mth.get('referencia'))) if on]
                    txt = ("antes da V4" if mth.get('pre') else mth.get('status', '')) + ("\n" + " · ".join(tags) if tags else "")
                    put(f'{hcols[j]}{st}', txt, font(8, bool(tags), GRAY), fill(LIGHT) if tags else None, al=HWRAP, bd=Border(bottom=thin), sheet=met)
            for i2, lin in enumerate(linhas_):
                label, fmt, vals = lin[0], lin[1], lin[2]; estilo = lin[3] if len(lin) > 3 else None
                r2 = st + 1 + i2
                put(f'B{r2}', label, font(9, estilo != 'sub', INK), al=LEFT, bd=Border(bottom=thin), sheet=met)
                for j, v in enumerate(vals):
                    put(f'{hcols[j]}{r2}', v, font(9, estilo == 'ating', GRAY), None, fmt, RIGHT, Border(bottom=thin), sheet=met)
                if estilo == 'ating' and hcols:   # atingimento: vermelho abaixo de 100%, verde a partir de 100%
                    faixa = f'C{r2}:{hcols[-1]}{r2}'
                    met.conditional_formatting.add(faixa, CellIsRule(operator='lessThan', formula=['1'], font=Font(name=FONT, size=9, bold=True, color=RED)))
                    met.conditional_formatting.add(faixa, CellIsRule(operator='greaterThanOrEqual', formula=['1'], font=Font(name=FONT, size=9, bold=True, color=GREEN)))
            return st + len(linhas_) + 2
        if HIST:
            rr = tabela_meses("HISTÓRICO · MESES LIDOS DA FONTE",
                              HIST.get('nota', "Meses marcados como janela alimentam as taxas atuais; 'referência' é a evidência do nível atingível."),
                              HIST['meses'], HIST['linhas'], rr)
        if LEG:
            rr = tabela_meses(LEG.get('titulo', "PROJETADO × REALIZADO · O QUE A PROJEÇÃO ANTERIOR PROMETEU"),
                              LEG.get('nota', "O que uma projeção anterior prometia mês a mês, contra o que a fonte registra como realizado. "
                                              "Atingimento = realizado ÷ prometido; fica em branco quando falta um dos dois."),
                              LEG['meses'], LEG['linhas'], rr)
        if ENV:
            put(f'B{rr}', "ENVELOPE HISTÓRICO E ALVO DA RAMPA", font(9, True, RED), sheet=met)
            put(f'B{rr+1}', ENV.get('nota', ''), font(9, False, GRAY, True), sheet=met)
            eh = rr + 2; met.row_dimensions[eh].height = 34
            heads = ["ALAVANCA", "ATUAL (JANELA)", "MEDIANA DO PERÍODO", "MELHOR MÊS", "QUANDO", "MÊS DE REFERÊNCIA", f"ALVO DA RAMPA ({ENV.get('alvo_em', '')})"]
            for c2, h2 in zip(['B', 'C', 'D', 'E', 'F', 'G', 'H'], heads):
                put(f'{c2}{eh}', h2, font(8, True, WHITE), fill(BLACK), al=HWRAP if c2 != 'B' else HWRAPL, sheet=met)
            for i2, (label, fmt, atual, mediana, melhor, melhor_mes, referencia, alvo) in enumerate(ENV['linhas']):
                r2 = eh + 1 + i2
                put(f'B{r2}', label, font(9, True, INK), al=LEFT, bd=Border(bottom=thin), sheet=met)
                for c2, v in zip(['C', 'D', 'E', 'G'], [atual, mediana, melhor, referencia]):
                    put(f'{c2}{r2}', v, font(9, False, GRAY), None, fmt, RIGHT, Border(bottom=thin), sheet=met)
                put(f'F{r2}', melhor_mes, font(8, False, GRAY), al=CENTER, bd=Border(bottom=thin), sheet=met)
                put(f'H{r2}', alvo, font(9, True, INK), fill(LIGHT), fmt, RIGHT, Border(bottom=thin), sheet=met)
            rr = eh + len(ENV['linhas']) + 2
        if BASE:
            put(f'B{rr}', BASE.get('titulo', "CENÁRIO DE REFERÊNCIA (VALORES DO PILOTO)"), font(9, True, RED), sheet=met)
            put(f'B{rr+1}', BASE.get('nota'), font(9, False, GRAY, True), sheet=met)
            hdr = rr + 2; met.row_dimensions[hdr].height = 20
            put(f'B{hdr}', "MÉTRICA", font(9, True, WHITE), fill(BLACK), al=HWRAPL, sheet=met)
            for k, col in enumerate(SUPC): put(f'{col}{hdr}', f"M{k+1}", font(8, True, WHITE), fill(BLACK), al=HWRAP, sheet=met)
            for i, (label, fmt, vals) in enumerate(BASE['linhas']):
                r2 = hdr + 1 + i
                put(f'B{r2}', label, font(9, True, INK), al=LEFT, bd=Border(bottom=thin), sheet=met)
                for k, col in enumerate(SUPC):
                    put(f'{col}{r2}', vals[k] if k < len(vals) else None, font(9, False, GRAY), None, fmt, RIGHT, Border(bottom=thin), sheet=met)
            rr = hdr + len(BASE['linhas']) + 2
        rr = secoes(cfg.get('metodologia_fim'), rr)   # fontes e notas que ficam no fim da aba
        ult = get_column_letter(max(14, 2 + max([len(SUPC)] + [len(x['meses']) for x in (HIST, LEG) if x])))
        met.print_area = f"A1:{ult}{rr}"
        met.page_setup.orientation = 'landscape'; met.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        met.page_setup.fitToWidth = 1; met.page_setup.fitToHeight = 0

    # ---- visual / impressão
    ws.freeze_panes = 'C1'
    ws.print_area = f"A1:{max(TOT_R, KPI_LAST)}{LAST_ROW}"
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = cfg['footer']; ws.oddFooter.center.size = 8
    return {'R': R, 'M': M, 'P': P, 'HDR1': HDR1, 'HDR2': HDR2, 'META_TOP': META_TOP, 'LAST_TABLE': LAST_TABLE, 'LAST_ROW': LAST_ROW, 'support': sup.title}


# ===================================================================== blocos comuns
def kpi_payback_cards(ctx, cum_key='cum', flag_key='flag'):
    cum, flag = ctx.rng(cum_key), ctx.rng(flag_key)
    return {
      'payback': ("PAYBACK", f'=IF(COUNTIF({flag},1)=0,"Não atinge no período","Mês "&MATCH(1,{flag},0))', None, "1º mês com acumulado ≥ 0"),
      'exposicao': ("MAIOR EXPOSIÇÃO DE CAIXA", f"=MAX(0,-MIN({cum}))", FMT_BRL, f'=IF(MIN({cum})>=0,"acumulado nunca fica negativo","em Mês "&MATCH(MIN({cum}),{cum},0))'),
      'azul': ("NO AZUL A PARTIR DE", f'=IF(COUNTIF({ctx.rng("flag_cont")},1)=0,"Não acontece no período","Mês "&MATCH(1,{ctx.rng("flag_cont")},0))', None, "daí em diante, todo mês com resultado ≥ 0"),
      'excedente': ("EXCEDENTE PÓS-PAYBACK", f'=IF(COUNTIF({flag},1)=0,"—",{ctx.last(cum_key)}-INDEX({cum},1,MATCH(1,{flag},0)))', FMT_BRL, "gerado do payback até o fim do período"),
    }

def meta_block(ctx, m, target_default, receita_key, ticket_key, vendas_key, custo_key, cum_key, margem_key=None, comissao_key=None, acum0_key=None):
    """Engenharia reversa do mês-alvo, coerente com a tabela (aplica margem e, se houver, comissão)."""
    cum = ctx.rng(cum_key)
    denom = f"INDEX({ctx.rng(margem_key)},1,{m[0]})" if margem_key else "1"
    if comissao_key: denom = f"(INDEX({ctx.rng(comissao_key)},1,{m[0]})*{denom})"
    receita_label = "GMV necessário no mês-alvo" if comissao_key else ("Receita faturada necessária no mês-alvo" if margem_key else "Receita necessária no mês-alvo")
    acum0 = ctx.prem(acum0_key) if acum0_key else None
    base0 = acum0 if acum0 else "0"
    return [
      (f"Mês-alvo do breakeven (1 a {ctx.n})", target_default, '0', True),
      ("Acumulado até o mês anterior", f"=IF({m[0]}<=1,{base0},INDEX({cum},1,{m[0]}-1))", FMT_BRL, False),
      ("Custo do mês-alvo (fee + mídia)", f"=INDEX({ctx.rng(custo_key)},1,{m[0]})", FMT_BRL, False),
      ("Resultado MC necessário no mês-alvo", f"=MAX(0,{m[2]}-{m[1]})", FMT_BRL, False),
      (receita_label, f"=IFERROR({m[3]}/{denom},0)", FMT_BRL, False),
      ("Vendas necessárias (arredondado p/ cima)", f"=IFERROR(ROUNDUP({m[4]}/INDEX({ctx.rng(ticket_key)},1,{m[0]}),0),0)", FMT_INT, False),
      ("Vendas projetadas no mês-alvo", f"=INDEX({ctx.rng(vendas_key)},1,{m[0]})", FMT_INT, False),
      ("Gap de vendas (necessárias − projetadas)", f"={m[5]}-{m[6]}", FMT_INT, False),
    ]
